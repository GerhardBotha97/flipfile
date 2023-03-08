import os
import re
import requests
from fpdf import FPDF
from docx import Document
from openpyxl import Workbook

#############################################################
################## Make the files go brrrr ##################
#############################################################

# Define the string to search for and the replacement variable
search_string = "changemeplease"
replacement_string = "Hello there! :D"

# Define the input and output directories
input_directory = './templates'
output_directory = './output'

# List of URLs of images to download
urls = [
    # "https://example.com/image1.jpg",
    # "https://example.com/image2.jpeg",
    # "https://example.com/image3.gif",
    # "https://example.com/image4.png",
    # Add more URLs here
]

# Loop through each file in the input directory
for filename in os.listdir(input_directory):

    # Read the contents of the input file
    input_file_path = os.path.join(input_directory, filename)
    with open(input_file_path, "r") as input_file:
        file_contents = input_file.read()

    # Replace the search string with the replacement string using regular expressions
    modified_file_contents = re.sub(search_string, replacement_string, file_contents)

    # Write the modified contents to the output file
    output_file_path = os.path.join(output_directory, filename)
    with open(output_file_path, "w") as output_file:
        output_file.write(modified_file_contents)

#############################################################
################## Fetch those images mah dude ##################
#############################################################

# Specify the directory where you want to save the images
save_dir = output_directory

# Loop over each URL and download the corresponding image
for i, url in enumerate(urls):
    # Send a request to the URL to get the image data
    response = requests.get(url)

    # Check if the request was successful
    if response.status_code == 200:
        # Determine the file extension of the image
        ext = url.split(".")[-1].lower()

        # Construct the filename to use for the saved image
        filename = f"file_{i+1}.{ext}"

        # Construct the full path where the image will be saved
        save_path = os.path.join(save_dir, filename)

        # Save the image data to a file
        with open(save_path, "wb") as f:
            f.write(response.content)

        print(f"Image {i+1} saved successfully.")

    else:
        print(f"Failed to download image {i+1}.")

#########################################################
################## Lets craft a PDF ####################
#########################################################

# Create a new PDF object
pdf = FPDF()

# Add a page
pdf.add_page()

# Set the font and font size
pdf.set_font('Arial', 'B', 16)

# Write some text
pdf.cell(40, 10, '{}'.format(replacement_string))

# Save the PDF file
pdf.output('./output/file.pdf', 'F')

#########################################################
################## Lets make some XLSX ##################
#########################################################

# create a new workbook
workbook = Workbook()

# select the active worksheet
worksheet = workbook.active

# write some data to the worksheet
worksheet["A1"] = "{}".format(replacement_string)

# save the workbook
workbook.save("./output/file.xlsx")

#########################################################
################## Lets make some DOCX ##################
#########################################################

doc = Document()

# Add a paragraph to the document
doc.add_paragraph('{}'.format(replacement_string))

# Save the document
doc.save('./output/file.docx')

#########################################################
################## Lets compile that C ##################
#########################################################

# Needs to compile into EXE, DLL, ELF

# le EXE
os.system('x86_64-w64-mingw32-gcc ./output/file.c -o ./output/file.exe')

# duh DLL
os.system('x86_64-w64-mingw32-gcc -shared -o ./output/file.dll ./output/file.c')


